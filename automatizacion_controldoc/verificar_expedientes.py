"""
Verifica en ControlDoc que cada radicado tenga el expediente, la serie y la
subserie que dice la base (pestañas EXPEDIENTES 25 y EXPEDIENTES 26).

Fase 1: SOLO LECTURA. No modifica nada en ControlDoc. Para los radicados que
necesitan corrección deja el navegador abierto en Documentos > Actualizar
Tipología (si se usa --pausar) y escribe todo en un Excel de resultados.

Uso:
    python verificar_expedientes.py --base "BASE DILIGENCIAS 2026.xlsx" \
        --radicados DocumentosExpediente.xlsx [--pausar] [--limite 3]

--radicados acepta un .xlsx exportado de ControlDoc (columna RADICADO) o un
.txt con un radicado por línea.
"""
import argparse
import re
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

import openpyxl

URL = "https://sgdea.personeriabogota.gov.co/ControlDoc/Home/Index#"
HOJAS_BASE = ["EXPEDIENTES 26", "EXPEDIENTES 25"]
PATRON_RADICADO = re.compile(r"[ERI]-20\d\d-\d{7}")

# Serie de la base -> nombres aceptados en ControlDoc (normalizados).
# Ajustar si la serie/subserie correcta en ControlDoc se llama distinto.
EQUIVALENCIAS_SERIE = {
    "PROCESOS (INTERVENCION)": ["PROCESOS DE INTERVENCION", "PROCESOS (INTERVENCION)", "PROCESO DE INTERVENCION"],
    "PROCESO DE INTERVENCION": ["PROCESOS DE INTERVENCION", "PROCESOS (INTERVENCION)", "PROCESO DE INTERVENCION"],
    "DERECHOS DE PETICION": ["DERECHOS DE PETICION"],
    "DERECHO DE PETICION": ["DERECHOS DE PETICION"],
}

# Selectores de la página. Si ControlDoc cambia o alguno no funciona, solo se
# ajusta aquí.
SEL = {
    # ícono de la lupa del encabezado; si no se encuentra, el script pide hacer el clic a mano
    "lupa": "[title*='squeda' i], [data-original-title*='squeda' i], [title*='buscar' i]",
    "campo_radicado": "xpath=//*[normalize-space(text())='Radicado']/following::input[1]",
    "boton_buscar": "button:has-text('Buscar')",
    "modal_info": "text=INFORMACIÓN GENERAL",
    "boton_documentos": "button:has-text('Documentos'), a:has-text('Documentos')",
    "actualizar_tipologia": "text=Actualizar Tipología",
    "cerrar_modal": ".modal.show .close, .k-window-action, [aria-label='Close']",
}


def normalizar(texto):
    if texto is None:
        return ""
    texto = unicodedata.normalize("NFKD", str(texto))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", texto).strip().upper()


def cargar_base(ruta):
    """Devuelve {radicado: {hoja, fila, serie, clase, expediente}}."""
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    indice = {}
    for hoja in HOJAS_BASE:
        if hoja not in wb.sheetnames:
            print(f"Aviso: la base no tiene la pestaña {hoja}")
            continue
        for n, fila in enumerate(wb[hoja].iter_rows(min_row=3, values_only=True), start=3):
            fila = list(fila) + [None] * 13
            # columnas: C=RADICADO, D=SERIE, E=CLASE DE EXPEDIENTE, L=No. EXPEDIENTE
            celdas = " ".join(str(x) for x in fila[1:3] if x)
            for rad in PATRON_RADICADO.findall(celdas):
                indice.setdefault(rad, {
                    "hoja": hoja, "fila": n,
                    "serie": normalizar(fila[3]), "clase": normalizar(fila[4]),
                    "expediente": str(fila[11]).strip() if fila[11] not in (None, "") else "",
                })
    return indice


def cargar_radicados(ruta):
    """Lista de (radicado, asunto). El asunto sirve para hallar el radicado al que responde."""
    ruta = Path(ruta)
    if ruta.suffix.lower() == ".txt":
        return [(l.strip(), "") for l in ruta.read_text(encoding="utf-8").splitlines() if l.strip()]
    ws = openpyxl.load_workbook(ruta, read_only=True, data_only=True).active
    filas = ws.iter_rows(values_only=True)
    encabezado = [normalizar(re.sub(r"<[^>]+>", "", str(c or ""))) for c in next(filas)]
    i_rad = encabezado.index("RADICADO")
    i_asu = encabezado.index("ASUNTO") if "ASUNTO" in encabezado else None
    salida = []
    for f in filas:
        if f[i_rad]:
            salida.append((str(f[i_rad]).strip(), str(f[i_asu] or "") if i_asu is not None else ""))
    return salida


def buscar_en_base(indice, radicado, asunto):
    """Busca el radicado; si no está, prueba con los radicados citados en el asunto."""
    if radicado in indice:
        return radicado, indice[radicado]
    for ref in PATRON_RADICADO.findall(asunto or ""):
        if ref in indice:
            return ref, indice[ref]
    return None, None


def serie_corresponde(serie_base, serie_cd, subserie_cd):
    if not serie_base:
        return None
    aceptadas = EQUIVALENCIAS_SERIE.get(serie_base, [serie_base])
    return any(a in serie_cd or a in subserie_cd for a in aceptadas)


def leer_info(page):
    """Lee la tabla de INFORMACIÓN GENERAL como {ETIQUETA: valor}."""
    page.wait_for_selector(SEL["modal_info"], timeout=20000)
    page.wait_for_timeout(800)
    pares = page.evaluate("""() => {
        const out = {};
        document.querySelectorAll('tr').forEach(tr => {
            const td = tr.querySelectorAll('td, th');
            if (td.length >= 2) out[td[0].innerText.trim().toUpperCase()] = td[1].innerText.trim();
        });
        return out;
    }""")
    return {normalizar(k): v for k, v in pares.items()}


def abrir_buscador(page):
    page.goto(URL)
    page.wait_for_load_state("networkidle")
    campo = page.locator(SEL["campo_radicado"]).first
    if not (campo.count() and campo.is_visible()):
        try:
            page.locator(SEL["lupa"]).first.click(timeout=5000)
        except Exception:
            input("No encontré la lupa. Haga clic en ella en el navegador y presione ENTER aquí... ")
        campo.wait_for(state="visible", timeout=20000)


def consultar(page, radicado):
    campo = page.locator(SEL["campo_radicado"]).first
    campo.fill("")
    campo.fill(radicado)
    page.locator(SEL["boton_buscar"]).last.click()
    fila = page.locator(f"tr:has(td:text-is('{radicado}'))").first
    fila.wait_for(timeout=20000)
    # primera columna: ícono "i" de INFORMACIÓN
    fila.locator("td").first.locator("a, i, img, span, button").first.click()
    return leer_info(page)


def cerrar_info(page):
    page.keyboard.press("Escape")
    cerrar = page.locator(SEL["cerrar_modal"])
    if cerrar.count():
        try:
            cerrar.last.click(timeout=2000)
        except Exception:
            pass


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", required=True)
    ap.add_argument("--radicados", required=True)
    ap.add_argument("--salida", default=f"resultado_{datetime.now():%Y%m%d_%H%M}.xlsx")
    ap.add_argument("--perfil", default="perfil_chrome", help="carpeta donde se guarda la sesión del navegador")
    ap.add_argument("--limite", type=int, default=0, help="procesar solo los primeros N (para pruebas)")
    ap.add_argument("--pausar", action="store_true",
                    help="en los que requieren corrección, abrir Actualizar Tipología y esperar")
    args = ap.parse_args()

    indice = cargar_base(args.base)
    radicados = cargar_radicados(args.radicados)
    if args.limite:
        radicados = radicados[: args.limite]
    print(f"Base: {len(indice)} radicados indexados. A revisar: {len(radicados)}")

    from playwright.sync_api import sync_playwright

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resultado"
    ws.append(["RADICADO", "RADICADO EN BASE", "HOJA/FILA BASE", "EXPEDIENTE BASE", "SERIE BASE",
               "EXPEDIENTE CONTROLDOC", "SERIE CONTROLDOC", "SUBSERIE CONTROLDOC", "TIPOLOGIA CONTROLDOC",
               "EXPEDIENTE OK", "SERIE OK", "ACCION"])

    with sync_playwright() as p:
        ctx = p.chromium.launch_persistent_context(args.perfil, channel="chrome", headless=False)
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        page.goto(URL)
        input("Inicie sesión en ControlDoc en la ventana que se abrió y presione ENTER aquí... ")
        abrir_buscador(page)

        for radicado, asunto in radicados:
            ref, base = buscar_en_base(indice, radicado, asunto)
            try:
                info = consultar(page, radicado)
            except Exception as e:
                print(f"{radicado}: ERROR al consultar ({e.__class__.__name__})")
                ws.append([radicado, ref, "", "", "", "", "", "", "", "", "", f"ERROR: {e}"[:300]])
                wb.save(args.salida)
                abrir_buscador(page)
                continue

            exp_cd = (info.get("EXPEDIENTE(S)-ID") or "").strip()
            serie_cd = normalizar(info.get("SERIE"))
            sub_cd = normalizar(info.get("SUBSERIE"))
            tipo_cd = info.get("TIPOLOGIA DOCUMENTAL", "")
            exp_vacio = exp_cd in ("", "0")

            if base is None:
                exp_ok, serie_ok = None, None
                accion = "NO ESTA EN LA BASE - revisar manualmente"
            else:
                exp_ok = (not exp_vacio) and base["expediente"] != "" and base["expediente"] in re.split(r"[^\d]+", exp_cd)
                serie_ok = serie_corresponde(base["serie"], serie_cd, sub_cd)
                if exp_ok and serie_ok:
                    accion = "OK"
                elif exp_vacio and not serie_ok:
                    accion = "ACTUALIZAR TIPOLOGIA (expediente vacío y serie no corresponde)"
                elif exp_vacio:
                    accion = "ASOCIAR EXPEDIENTE (serie correcta, expediente vacío)"
                elif not exp_ok:
                    accion = "REVISAR: expediente distinto al de la base"
                else:
                    accion = "REVISAR: serie no corresponde"

            print(f"{radicado}: {accion}")
            ws.append([radicado, ref, f"{base['hoja']} / fila {base['fila']}" if base else "",
                       base["expediente"] if base else "", base["serie"] if base else "",
                       exp_cd, serie_cd, sub_cd, tipo_cd,
                       {True: "SI", False: "NO", None: ""}[exp_ok], {True: "SI", False: "NO", None: ""}[serie_ok],
                       accion])
            wb.save(args.salida)

            if args.pausar and accion.startswith("ACTUALIZAR TIPOLOGIA"):
                page.locator(SEL["boton_documentos"]).first.click()
                page.locator(SEL["actualizar_tipologia"]).first.click()
                input(f"  {radicado}: abierto Actualizar Tipología (base: serie {base['serie']}, "
                      f"expediente {base['expediente']}). Haga el cambio y presione ENTER para seguir... ")
            cerrar_info(page)

        ctx.close()
    print(f"Listo. Resultados en {args.salida}")


if __name__ == "__main__":
    sys.exit(main())
